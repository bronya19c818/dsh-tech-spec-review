from __future__ import annotations

import hashlib
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .models import Block, Chapter, DocumentEvidence, ParagraphRole
from .numbering import NumberingResolver, attr

MANUAL = re.compile(r"^\s*(\d+(?:[.．]\d+)*)([.．、]?\s*)([^\d\s].*)$")
INLINE_DECIMAL_HEADING = re.compile(
    r"^\s*(?P<number>\d+(?:[.．]\d+)+)\s*"
    r"(?P<title>[^。！？!?；;\n]{1,40}?)(?P<separator>[。！？!?；;])"
    r"(?P<body>\s*\S.*)$"
)
CAPTION = re.compile(r"^\s*(表|图|附录)\s*([A-Z]|\d+(?:[.．-]\d+)*)\s*(?:[：:].*)?(.*)$")
PARAGRAPH_CAPTION = re.compile(
    r"^\s*(?P<kind>图|表)\s*(?P<number>[A-Z]|\d+(?:[.．-]\d+)*)"
    r"\s*(?:[：:]\s*|[、．.]\s*)?(?P<title>\S.*)$",
    re.IGNORECASE,
)
CHINESE_HEADING = re.compile(r"^([一二三四五六七八九十]+)([、．.])\s*(.+)$")
PARENTHESIZED_HEADING = re.compile(
    r"^(?:（(?P<full>[一二三四五六七八九十]+|\d+)）|"
    r"\((?P<half>[一二三四五六七八九十]+|\d+)\))\s*(?P<title>.+)$"
)
CIRCLED_HEADING = re.compile(r"^([①-⑳])\s*(.+)$")
RIGHT_PAREN_HEADING = re.compile(r"^(\d+)\)\s*(.+)$")
PART = re.compile(
    r"^(?:[（(][一二三四五六七八九十]+[）)]\s*)?(通用部分|专用部分)(?:格式如下[：:]?)?$"
)
APPENDIX_HEADING = re.compile(r"^附录\s*[A-Z](?:(?:\s*[、．.：:]\s*|\s+)\S.*)?$", re.IGNORECASE)
TOC_STYLE = re.compile(r"^(?:TOC(?:\s*(?:\d+|Heading))?|目录(?:\s*\d+)?|目录标题)$", re.IGNORECASE)
HEADING_STYLE = re.compile(r"^(?:Heading\s*\d+|标题\s*\d+|Heading)$", re.IGNORECASE)
TITLE_STYLE = re.compile(r"^(?:Title|Document Title|文档标题|封面标题|标题)$", re.IGNORECASE)
TOC_TITLE = re.compile(r"^(?:目录|Contents|Table of Contents)$", re.IGNORECASE)
TOC_ENTRY = re.compile(
    r"(?:\t|\.{2,}|…{2,}|·{2,})\s*[\dIVXLCDM]+(?:\s*/\s*\d+)?\s*$",
    re.IGNORECASE,
)
FRONT_MATTER_LABEL = re.compile(
    r"^\s*(?:项目名称|项目编号|招标编号|采购编号|采购人|招标人|物资名称|设备名称|"
    r"编制单位|编制日期|发布日期|版本|文件编号|技术规范书编号)\s*[：:]"
)
SPECIAL_NOTE_PREFIX = re.compile(r"^\s*(?:注(?:释)?|备注|注意)\s*[：:]")
DOCUMENT_TITLE = re.compile(r"技术规范书|技术规格书|技术标准|采购文件|招标文件")
NUMBERING_PREFIX = re.compile(
    r"^\s*(?P<marker>(?:\d+(?:[.．]\d+)*(?:[、.．)）])?|"
    r"[一二三四五六七八九十]+[、．.]|"
    r"[（(](?:[一二三四五六七八九十]+|\d+)[）)]|"
    r"[①-⑳]|[•●▪◦‣–—-]))\s*"
)
BODY_CUE = re.compile(
    r"^(?:是|为)|本文件|本规范|投标人|招标人|供应商|卖方|买方|应当|必须|不得|不应|"
    r"(?<![对适相供反响])应(?!急|用|力)|须(?!知|要)|(?<!供)需(?!求)|"
    r"应(?=满足|符合|提供|按|执行|达到|具备|具有|配置|保持|完成)|"
    r"需(?=提供|满足|符合|按|配置)|"
    r"提供|满足|确保|符合|包括|包含|执行|采用|配备|达到|不低于|不超过|按照|"
    r"参见|详见|见第|列出|给出|示出"
)
DECIMAL_EXPLANATION_CUE = re.compile(r"圆周率|近似值")
LIST_INTRO = re.compile(
    r"(?:包括|包含)(?:\s*(?:(?:以下|下列|如下)(?:所示)?[\u4e00-\u9fff0-9]{0,16}"
    r"|[\u4e00-\u9fff0-9]{0,8}))?$"
    r"|(?:以下|下列)[\u4e00-\u9fff0-9]{0,16}$"
    r"|(?:如下(?:所示)?|分为|分成|分别为|主要有)[\u4e00-\u9fff0-9]{0,16}$"
)


def chinese_number(text: str) -> str:
    digits = {char: i for i, char in enumerate("零一二三四五六七八九")}
    if "十" in text:
        a, b = text.split("十", 1)
        return str((digits.get(a, 1) * 10) + digits.get(b, 0))
    return str(digits[text])


def manual_identity(text: str, outline: int | None):
    if PART.match(text):
        return "", text, False, False, ""

    chinese = CHINESE_HEADING.match(text)
    if chinese and (outline is not None or len(chinese[3]) < 30):
        return chinese_number(chinese[1]), chinese[3], True, False, chinese[1] + chinese[2]

    parenthesized = PARENTHESIZED_HEADING.match(text)
    if parenthesized:
        raw_value = parenthesized["full"] or parenthesized["half"]
        token = f"（{raw_value}）" if parenthesized["full"] else f"({raw_value})"
        title = parenthesized["title"].strip()
        prose = re.search(r"本文件|本规范|应|必须|不得|投标人|招标人|卖方|买方|按照?", title)
        heading = outline is not None or (
            len(title) <= 40 and not prose and not re.search(r"[；;：:]", title)
        )
        number = str(int(raw_value)) if raw_value.isdecimal() else chinese_number(raw_value)
        return (
            number if heading else "",
            title if heading else text,
            heading,
            False,
            token if heading else "",
        )

    circled = CIRCLED_HEADING.match(text)
    if circled:
        title = circled[2].strip()
        prose = re.search(r"本文件|本规范|应|必须|不得|投标人|招标人|卖方|买方|按照?", title)
        heading = outline is not None or (
            len(title) <= 40 and not prose and not re.search(r"[；;：:]", title)
        )
        number = str(ord(circled[1]) - ord("①") + 1)
        return (
            number if heading else "",
            title if heading else text,
            heading,
            False,
            circled[1] if heading else "",
        )

    right_parenthesis = RIGHT_PAREN_HEADING.match(text)
    if right_parenthesis:
        title = right_parenthesis[2].strip()
        prose = re.search(r"本文件|本规范|应|必须|不得|投标人|招标人|卖方|买方|按照?", title)
        heading = outline is not None or (
            len(title) <= 40 and not prose and not re.search(r"[；;：:]", title)
        )
        return (
            str(int(right_parenthesis[1])) if heading else "",
            title if heading else text,
            heading,
            False,
            right_parenthesis[1] + ")" if heading else "",
        )

    match = MANUAL.match(text)
    if not match:
        return "", text, False, False, ""
    number, separator, body = match[1].replace("．", "."), match[2], match[3].strip()
    quantity = re.match(
        r"(?:个月|个|月|年|天|日|小时|分钟|秒|台|套|件|米|毫米|千克|吨|%|℃|k[VW]|M[Pp]a|mm|kg|m[²³]?)(?:内|外|以上|以下|之内|$|\s)",
        body,
    )
    if quantity or (outline is None and "." not in number and not separator):
        return "", text, False, False, ""
    # Limit this exception to explicit decimal-value explanations, not general titles using "表示".
    if (
        re.fullmatch(r"\d+\.\d{2,}", number)
        and body.startswith("表示")
        and DECIMAL_EXPLANATION_CUE.search(body)
    ):
        return "", text, False, False, ""
    title = body.split("。", 1)[0].strip()
    prose = re.search(r"本文件|本规范|应|必须|不得|投标人|招标人|卖方|买方|按[照照]?", title)
    heading = outline is not None or (
        len(title) <= 40 and not prose and not re.search(r"[；;：:]", title)
    )
    clause = not heading and "." in number
    return (
        number if heading or clause else "",
        title if heading else body,
        heading,
        clause,
        number + separator.rstrip(),
    )


def outline_level(paragraph) -> int | None:
    props = [paragraph._p.pPr]
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        props.append(style.element.pPr)
        style = style.base_style
    for pr in props:
        value = attr(pr, "outlineLvl")
        if value:
            return int(value) + 1 if int(value) < 9 else None
    return None


def _is_title_like(text: str) -> bool:
    value = text.strip()
    return (
        bool(value)
        and len(value) <= 80
        and not re.search(r"[。；;！？!?]", value)
        and not BODY_CUE.search(value)
    )


def _heading_text(text: str) -> tuple[str, str]:
    match = NUMBERING_PREFIX.match(text)
    if not match:
        return "", text.strip()
    return match["marker"].strip(), text[match.end() :].strip()


def _is_decimal_heading_number(number: str) -> bool:
    parts = re.split(r"[.．]", number.strip())
    return 2 <= len(parts) <= 4 and all(part.isdigit() and len(part) <= 3 for part in parts)


def _dotted_numeric_marker(marker: str) -> str | None:
    number = marker.strip().rstrip("、.．")
    return number if re.fullmatch(r"\d+(?:[.．]\d+)+", number) else None


def _is_list_marker(marker: str) -> bool:
    return bool(
        re.fullmatch(
            r"(?:\d+[)）]|\d+[、]|[（(]\d+[）)]|[①-⑳]|[•●▪◦‣–—-])",
            marker,
        )
    )


def _is_split_inline_heading(block: Block, blocks: list[Block], index: int) -> bool:
    return index + 1 < len(blocks) and blocks[index + 1].id == f"{block.id}-body"


def _is_structural_heading(block: Block, *, inline_heading: bool = False) -> bool:
    text = block.text.strip()
    if SPECIAL_NOTE_PREFIX.match(text):
        return False
    if PART.match(text) or APPENDIX_HEADING.match(text):
        return True
    marker, title = _heading_text(text)
    if not _is_title_like(title if marker else text):
        return inline_heading
    if _is_list_marker(marker):
        return False
    dotted_number = _dotted_numeric_marker(marker)
    if dotted_number:
        return _is_decimal_heading_number(dotted_number) and block.is_heading
    if block.outline_level is not None:
        return True
    heading_style = bool(HEADING_STYLE.match(block.style_name or ""))
    if not marker:
        return heading_style
    if re.fullmatch(r"[一二三四五六七八九十]+[、．.]", marker):
        return True
    return heading_style or bool(block.is_heading and re.fullmatch(r"\d+[.．]?", marker))


def _chinese_parenthesized_info(block: Block) -> tuple[int, str] | None:
    match = PARENTHESIZED_HEADING.match(block.text.strip())
    if not match:
        return None
    value = match["full"] or match["half"]
    title = match["title"].strip()
    if value.isdecimal() or not _is_title_like(title):
        return None
    return int(chinese_number(value)), title


def _is_chinese_parenthesized_candidate(block: Block) -> bool:
    return _chinese_parenthesized_info(block) is not None


def _is_list_intro(text: str) -> bool:
    value = re.sub(r"[：:。；;，,]+$", "", text.strip())
    return bool(LIST_INTRO.search(value))


def _has_parent_heading(blocks: list[Block], index: int, toc_ids: set[str]) -> bool:
    for previous_index in range(index - 1, -1, -1):
        previous = blocks[previous_index]
        if previous.kind != "paragraph" or previous.in_table or previous.id in toc_ids:
            continue
        if _is_chinese_parenthesized_candidate(previous):
            continue
        if _is_structural_heading(
            previous,
            inline_heading=_is_split_inline_heading(previous, blocks, previous_index),
        ):
            return True
    return False


def _parenthesized_heading_group_ids(blocks: list[Block], toc_ids: set[str]) -> set[str]:
    candidates = [
        (index, block, info[0])
        for index, block in enumerate(blocks)
        if block.id not in toc_ids and (info := _chinese_parenthesized_info(block)) is not None
    ]
    groups: list[list[tuple[int, Block, int]]] = []
    for candidate in candidates:
        if (
            groups
            and candidate[0] == groups[-1][-1][0] + 1
            and candidate[2] == groups[-1][-1][2] + 1
        ):
            groups[-1].append(candidate)
        else:
            groups.append([candidate])

    heading_ids: set[str] = set()
    for group in groups:
        first_index = group[0][0]
        preceding = next(
            (
                block
                for block in reversed(blocks[:first_index])
                if block.kind == "paragraph" and not block.in_table and block.id not in toc_ids
            ),
            None,
        )
        if any(block.outline_level is not None for _, block, _ in group):
            heading_ids.update(block.id for _, block, _ in group)
        elif preceding is not None and _is_list_intro(preceding.text):
            continue
        elif any(HEADING_STYLE.match(block.style_name or "") for _, block, _ in group):
            heading_ids.update(block.id for _, block, _ in group)
        elif _has_parent_heading(blocks, first_index, toc_ids):
            heading_ids.update(block.id for _, block, _ in group)
    return heading_ids


def _is_heading_role(
    block: Block,
    blocks: list[Block],
    index: int,
    toc_ids: set[str],
    parenthesized_heading_ids: set[str],
) -> bool:
    if block.in_table:
        return False
    text = block.text.strip()
    if FRONT_MATTER_LABEL.match(text) or SPECIAL_NOTE_PREFIX.match(text):
        return False

    if PART.match(text) or APPENDIX_HEADING.match(text):
        return True

    marker, title = _heading_text(text)
    inline_heading = _is_split_inline_heading(block, blocks, index)
    if not _is_title_like(title if marker else text):
        return inline_heading
    if _is_list_marker(marker):
        return False
    dotted_number = _dotted_numeric_marker(marker)
    if dotted_number:
        return _is_decimal_heading_number(dotted_number) and block.is_heading
    if block.outline_level is not None:
        return True

    heading_style = bool(HEADING_STYLE.match(block.style_name or ""))
    if not marker:
        return heading_style

    if re.fullmatch(r"[（(][一二三四五六七八九十]+[）)]", marker):
        return block.id in parenthesized_heading_ids
    # Outline-like decimal and Chinese chapter numbering provide structural
    # evidence; list markers such as 1), ①, and bullets remain body content.
    if re.fullmatch(r"[一二三四五六七八九十]+[、．.]", marker):
        return True
    if heading_style:
        return True
    return bool(block.is_heading and re.fullmatch(r"\d+[.．]?", marker))


def _split_inline_heading_blocks(blocks: list[Block]) -> tuple[list[Block], set[str]]:
    expanded: list[Block] = []
    body_ids: set[str] = set()
    for block in blocks:
        if block.in_table or not block.is_heading or not _is_decimal_heading_number(block.number):
            expanded.append(block)
            continue
        match = INLINE_DECIMAL_HEADING.match(block.text)
        if (
            match is None
            or not _is_decimal_heading_number(match["number"])
            or match["number"].replace("．", ".") != block.number
            or not _is_title_like(match["title"])
            or not match["body"].strip()
        ):
            expanded.append(block)
            continue
        heading_text = block.text[: match.end("separator")].strip()
        heading_number, heading_title, is_heading, _, _ = manual_identity(
            heading_text, block.outline_level
        )
        if not is_heading or heading_number.replace("．", ".") != block.number:
            expanded.append(block)
            continue
        body_id = f"{block.id}-body"
        expanded.append(block.model_copy(update={"text": heading_text, "title": heading_title}))
        expanded.append(
            block.model_copy(
                update={
                    "id": body_id,
                    "text": match["body"].strip(),
                    "source_text": block.text,
                    "number": "",
                    "title": match["body"].strip(),
                    "is_heading": False,
                    "is_clause": False,
                    "display_number": "",
                    "number_source": "",
                    "numbering": None,
                    "level": 0,
                    "outline_level": None,
                }
            )
        )
        body_ids.add(body_id)
    return expanded, body_ids


def _paragraph_toc_state(paragraph, field_stack: list[bool]) -> tuple[bool, bool]:
    """Track TOC fields across paragraph boundaries and identify page fields."""
    in_toc = False
    page_field = False
    for element in paragraph._p.iter():
        if element.tag == qn("w:fldChar"):
            field_type = element.get(qn("w:fldCharType"))
            if field_type == "begin":
                field_stack.append(False)
            elif field_type == "end" and field_stack:
                in_toc = in_toc or field_stack.pop()
        elif element.tag == qn("w:instrText"):
            instruction = element.text or ""
            if re.search(r"\bTOC\b", instruction, re.IGNORECASE):
                if field_stack:
                    field_stack[-1] = True
                in_toc = True
            if re.search(r"\b(?:PAGE|NUMPAGES)\b", instruction, re.IGNORECASE):
                page_field = True
        elif element.tag == qn("w:fldSimple"):
            instruction = element.get(qn("w:instr"), "")
            if re.search(r"\bTOC\b", instruction, re.IGNORECASE):
                in_toc = True
            if re.search(r"\b(?:PAGE|NUMPAGES)\b", instruction, re.IGNORECASE):
                page_field = True
        if any(field_stack):
            in_toc = True
    return in_toc, page_field


def _contains_visual(element) -> bool:
    if element is None:
        return False
    if element.tag == qn("w:tbl"):
        return True
    return any(
        node.tag in {qn("w:drawing"), qn("w:pict"), qn("w:object")} for node in element.iter()
    )


def _paragraph_has_nearby_visual(paragraph) -> bool:
    node = paragraph._p
    return (
        _contains_visual(node)
        or _contains_visual(node.getprevious())
        or _contains_visual(node.getnext())
    )


def _caption_role(text: str, nearby_visual: bool = False) -> ParagraphRole | None:
    match = PARAGRAPH_CAPTION.match(text)
    if not match:
        return None
    title = match["title"].strip()
    if re.match(r"^(?:为|是|中|内|的|应|仅|所|给出|列出|包含|示出|表示|如下|见)", title):
        return None
    if not _is_title_like(title) and not nearby_visual:
        return None
    return "figure_caption" if match["kind"] == "图" else "table_caption"


def _toc_blocks(blocks: list[Block]) -> set[str]:
    toc_ids: set[str] = set()
    entry_indexes: list[int] = []
    for index, block in enumerate(blocks):
        if block.kind != "paragraph" or block.in_table:
            continue
        if (
            block.toc_field
            or TOC_STYLE.match(block.style_name or "")
            or TOC_TITLE.match(block.text.strip())
        ):
            toc_ids.add(block.id)
        if TOC_ENTRY.search(block.text):
            entry_indexes.append(index)

    entry_set = set(entry_indexes)
    for index in entry_indexes:
        neighbors = [neighbor for neighbor in (index - 1, index + 1) if neighbor in entry_set]
        if neighbors:
            toc_ids.add(blocks[index].id)
            toc_ids.update(blocks[neighbor].id for neighbor in neighbors)
        # A single entry immediately following a contents heading is also a TOC entry.
        previous = index - 1
        if previous >= 0 and TOC_TITLE.match(blocks[previous].text.strip()):
            toc_ids.add(blocks[index].id)
            toc_ids.add(blocks[previous].id)
    return toc_ids


def _is_front_matter(block: Block, first_heading_index: int | None) -> bool:
    if block.in_table or block.paragraph_index is None:
        return False
    text = block.text.strip()
    if FRONT_MATTER_LABEL.match(text):
        return True
    if first_heading_index is not None and block.paragraph_index >= first_heading_index:
        return False
    title_style = bool(TITLE_STYLE.match(block.style_name or ""))
    if title_style and _is_title_like(text):
        return True
    return block.paragraph_index <= 2 and bool(DOCUMENT_TITLE.search(text)) and _is_title_like(text)


def _classify_roles(blocks: list[Block]) -> None:
    toc_ids = _toc_blocks(blocks)
    parenthesized_heading_ids = _parenthesized_heading_group_ids(blocks, toc_ids)
    heading_ids = {
        block.id
        for index, block in enumerate(blocks)
        if block.id not in toc_ids
        and _is_heading_role(block, blocks, index, toc_ids, parenthesized_heading_ids)
    }
    heading_indexes = [
        block.paragraph_index
        for block in blocks
        if block.id in heading_ids and block.paragraph_index is not None
    ]
    first_heading_index = min(heading_indexes, default=None)

    for block in blocks:
        if block.in_table:
            block.role = "table_content"
        elif block.id in toc_ids:
            block.role = "toc"
        else:
            caption_role = _caption_role(block.text, nearby_visual=block.nearby_visual)
            if caption_role:
                block.role = caption_role
            elif SPECIAL_NOTE_PREFIX.match(block.text.strip()):
                block.role = "other_special"
            elif block.id in heading_ids:
                block.role = "heading"
            elif _is_front_matter(block, first_heading_index):
                block.role = "front_matter"
            elif block.page_field or re.fullmatch(
                r"(?:第\s*)?\d+(?:\s*/\s*\d+)?(?:\s*页)?", block.text.strip()
            ):
                block.role = "other_special"
            else:
                block.role = "body"


def parse_docx(path: Path | str) -> DocumentEvidence:
    path = Path(path).resolve()
    document = Document(str(path))
    numbering = NumberingResolver(document)
    blocks: list[Block] = []
    part = "全文"
    paragraph_index = 0
    top_level_paragraph_index = 0
    table_index = 0
    toc_field_stack: list[bool] = []

    def add_paragraph(paragraph, block_id: str, kind: str):
        nonlocal part, paragraph_index
        paragraph_index += 1
        toc_field, page_field = _paragraph_toc_state(paragraph, toc_field_stack)
        text = paragraph.text.strip()
        if not text:
            return
        part_marker = PART.match(text)
        if part_marker:
            part = part_marker[1]
        visible, number, auto_level = numbering.resolve(paragraph)
        level = outline_level(paragraph)
        title = text
        manual, manual_title, manual_heading, clause, display = manual_identity(text, level)
        source = "word_auto_numbering" if visible else ""
        if manual:
            number, title, source = manual, manual_title, "manual"
        elif visible:
            text = f"{visible} {text}"
        heading = kind == "paragraph" and bool(number) and (level is not None or manual_heading)
        runs = [r for r in paragraph.runs if r.text.strip()]
        style_name = paragraph.style.name if paragraph.style else None
        numbering_text = visible.strip() if visible else ""
        if not numbering_text:
            numbering_match = NUMBERING_PREFIX.match(text)
            numbering_text = numbering_match["marker"].strip() if numbering_match else ""
        nearby_visual = _paragraph_has_nearby_visual(paragraph)
        fmt = {
            "style": style_name,
            "outline_level": level,
            "auto_level": auto_level,
        }
        if runs:
            run = runs[0]
            fmt.update(
                {
                    "font": run.font.name,
                    "size_pt": run.font.size.pt if run.font.size else None,
                    "bold": run.font.bold,
                }
            )
        blocks.append(
            Block(
                id=block_id,
                text=text,
                paragraph_index=paragraph_index,
                style_name=style_name,
                outline_level=level,
                numbering=numbering_text or None,
                in_table=kind == "cell",
                toc_field=toc_field,
                page_field=page_field,
                nearby_visual=nearby_visual,
                kind=kind,
                part=part,
                number=number,
                title=title,
                is_heading=heading,
                is_clause=clause and kind == "paragraph",
                display_number=display if manual else visible,
                number_source=source,
                level=number.count(".") + 1 if heading else 0,
                format=fmt,
            )
        )

    def walk_table(table, prefix: str):
        for ri, row in enumerate(table._tbl.tr_lst, 1):
            for ci, tc in enumerate(row.tc_lst, 1):
                cell = _Cell(tc, table)
                for pi, child in enumerate(tc, 1):
                    bid = f"{prefix}r{ri:03d}c{ci:03d}p{pi:03d}"
                    if child.tag == qn("w:p"):
                        add_paragraph(Paragraph(child, cell), bid, "cell")
                    elif child.tag == qn("w:tbl"):
                        walk_table(Table(child, cell), bid + "t")

    for child in document.element.body:
        if child.tag == qn("w:p"):
            top_level_paragraph_index += 1
            add_paragraph(
                Paragraph(child, document),
                f"p{top_level_paragraph_index:04d}",
                "paragraph",
            )
        elif child.tag == qn("w:tbl"):
            table_index += 1
            walk_table(Table(child, document), f"t{table_index:03d}")

    blocks, inline_body_ids = _split_inline_heading_blocks(blocks)
    _classify_roles(blocks)
    for block in blocks:
        if block.id in inline_body_ids:
            block.role = "body"

    chapters: list[Chapter] = []
    current = None
    for block in blocks:
        if block.role == "toc":
            continue
        if block.role == "heading" and block.number:
            current = Chapter(
                id=f"CH{len(chapters) + 1:03d}",
                part=block.part,
                number=block.number,
                title=block.title,
            )
            chapters.append(current)
        if current is not None:
            current.evidence_ids.append(block.id)
    sections = []
    for section in document.sections:
        sections.append(
            {
                key: round(getattr(section, key).mm, 2)
                if getattr(section, key) is not None
                else None
                for key in (
                    "page_width",
                    "page_height",
                    "top_margin",
                    "bottom_margin",
                    "left_margin",
                    "right_margin",
                )
            }
        )
    limitations = list(dict.fromkeys(numbering.limitations))
    if document.element.findall(".//" + qn("w:txbxContent")):
        limitations.append("文本框内容未纳入正文证据，请单独复核。")
    if not chapters:
        limitations.append("未识别到编号章节，全文作为一个内容审查单元。")
        chapters = [
            Chapter(
                id="CH001",
                part="全文",
                number="",
                title="全文",
                evidence_ids=[b.id for b in blocks],
            )
        ]
    return DocumentEvidence(
        source=str(path),
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        title=blocks[0].text if blocks else path.stem,
        blocks=blocks,
        chapters=chapters,
        sections=sections,
        limitations=limitations,
    )
